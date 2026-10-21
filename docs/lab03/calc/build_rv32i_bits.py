#!/usr/bin/env python3
"""LAB-03 RV32I bit-field atlas for the actual pinned PicoRV32 ROM.

Adds three instruction-decoding worksheets to the existing Excel Observatory.
Data provenance: LAB-03 boot ROM SHA-256 and embedded verified T3 observations.
It does not change, reinterpret or invent a physical measurement.
Uses openpyxl (no LibreOffice automation) and retains cached old formula values.
"""
from __future__ import annotations
import argparse
import collections
import copy
import hashlib
import json
import re
import shutil
import tempfile
import xml.etree.ElementTree as ET
from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED

import openpyxl
from openpyxl import load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.utils import get_column_letter
from openpyxl.workbook.properties import CalcProperties

ROM_SHA256 = '14b85bc20ce43cdf4ce09ef1db7a20af5677c1c728c64b71258f6c63ab5b8a3d'
DARK='102039'; PANEL='172B47'; TEXT='EEF4FF'; MUTED='AFC0D7'; TEAL='00A9A5'; CYAN='D7F1F0'; PURPLE='7257B4'; GOLD='FFF0CD'; BLUE='175AC4'; RED='FBE2DA'; GREEN='157552'; WHITE='FFFFFF'; LIGHT='F5F9FF'; GRAY='65768D'; DIVIDER='DBE6EF'
REGS=['zero','ra','sp','gp','tp','t0','t1','t2','s0','s1','a0','a1','a2','a3','a4','a5','a6','a7','s2','s3','s4','s5','s6','s7','s8','s9','s10','s11','t3','t4','t5','t6']
INFO={
 'addi':('Add Immediate','I','Сложить значение rs1 со знаковой 12-битной константой и записать в rd.','Подготавливает UART-символы, маски GPIO и промежуточные значения.'),
 'andi':('AND Immediate','I','Побитовое И значения rs1 со знаковой 12-битной константой.','Отделяет бит готовности при опросе UART.'),
 'beq':('Branch if Equal','B','Если значения двух регистров равны, изменить PC на заданное смещение.','Проверяет команду или организует цикл ожидания UART.'),
 'jal':('Jump And Link','J','Записать PC+4 в rd и перейти по PC-относительному смещению.','Вызывает подпрограммы; jal zero,... делает переход без сохранения адреса возврата.'),
 'lw':('Load Word','I','Загрузить 32-битное слово из памяти/MMIO по rs1+imm в rd.','Читает регистры UART и LED; FETCH lw не является наблюдением DATA_READ.'),
 'sw':('Store Word','S','Записать 32-битное значение rs2 по адресу rs1+imm.','PC=0x100 пишет выбранную LED-маску в MMIO 0x20000000.'),
 'lui':('Load Upper Immediate','U','Поместить 20-битный immediate в старшие биты rd; нижние 12 бит обнулить.','Формирует в t0 базовый адрес MMIO 0x20000000.'),
 'bne':('Branch if Not Equal','B','Если значения двух регистров различны, изменить PC на заданное смещение.','Проверяет результат обратного чтения LED-регистра.'),
 'srli':('Shift Right Logical Immediate','I-shift','Логически сдвинуть 32-битное значение rs1 вправо на shamt.','Выделяет шестнадцатеричные цифры для печати T3.'),
 'jalr':('Jump And Link Register','I','Записать PC+4 в rd; перейти на (rs1+imm) с обнулённым младшим битом.','Используется для возврата из функций печати UART.'),
}

# Type-conforming colors for a conventional RISC-V instruction map.
COLORS={'imm':('E7EEF8','415675'), 'rs1':('DDF4F0','126856'), 'rs2':('DDF4F0','126856'), 'rd':('DDF4F0','126856'), 'funct':('EFE8FF','64439E'), 'opcode':('FFF0CD','946518'), 'shamt':('DFEEFE','1F619F')}

def sg(value:int,width:int)->int:
    return value-(1<<width) if value & (1<<(width-1)) else value

def field(msb:int,lsb:int,label:str,w:int,meaning:str):
    bitstr=f'{(w>>lsb)&((1<<(msb-lsb+1))-1):0{msb-lsb+1}b}'
    return {'range':str(msb) if msb==lsb else f'{msb}–{lsb}', 'bits':bitstr,'label':label,'meaning':meaning,
            'n':msb-lsb+1,'start':lsb,'end':msb}

def decode(w:int,pc:int):
    op=w&127; rd=(w>>7)&31; f3=(w>>12)&7; rs1=(w>>15)&31; rs2=(w>>20)&31
    name='';fmt='';imm=None;target=None;hint=''
    if op==0x37:
        name='lui';fmt='U';imm=w>>12
        asm=f'lui {REGS[rd]}, 0x{(w&0xfffff000):08x}'
        fs=[field(31,12,'imm[31:12]',w,'20-битная константа; значение в rd = imm << 12'),field(11,7,'rd',w,f'Приёмник x{rd} = {REGS[rd]}'),field(6,0,'opcode',w,'0110111: LUI')]
        hint=f'rd ({REGS[rd]}) ← 0x{w&0xfffff000:08X}. Здесь формируется старшая часть адреса периферии.'
    elif op==0x6F:
        name='jal';fmt='J';raw=((w>>31&1)<<20)|((w>>12&255)<<12)|((w>>20&1)<<11)|((w>>21&1023)<<1)
        imm=sg(raw,21);target=(pc+imm)&0xffffffff
        asm=f'jal {REGS[rd]}, 0x{target:08x}'
        fs=[field(31,31,'imm[20]',w,'Знаковый бит PC-относительного перехода'),field(30,21,'imm[10:1]',w,'Средние биты адресного смещения'),field(20,20,'imm[11]',w,'Переставленный бит смещения'),field(19,12,'imm[19:12]',w,'Биты смещения перехода'),field(11,7,'rd',w,f'Приёмник адреса возврата: x{rd} = {REGS[rd]}'),field(6,0,'opcode',w,'1101111: JAL')]
        hint=f'PC ← 0x{target:08X}; rd ({REGS[rd]}) ← PC+4. Смещение {imm:+d} байт относительно PC.'
    elif op==0x63:
        if f3 not in (0,1):raise ValueError('unsupported branch funct3')
        name={0:'beq',1:'bne'}[f3];fmt='B'
        raw=((w>>31&1)<<12)|((w>>7&1)<<11)|((w>>25&63)<<5)|((w>>8&15)<<1)
        imm=sg(raw,13);target=(pc+imm)&0xffffffff
        asm=f'{name} {REGS[rs1]}, {REGS[rs2]}, 0x{target:08x}'
        fs=[field(31,31,'imm[12]',w,'Знаковый бит ветвления'),field(30,25,'imm[10:5]',w,'Старшие биты смещения'),field(24,20,'rs2',w,f'Второй сравниваемый регистр x{rs2} = {REGS[rs2]}'),field(19,15,'rs1',w,f'Первый сравниваемый регистр x{rs1} = {REGS[rs1]}'),field(14,12,'funct3',w,'000=BEQ, 001=BNE'),field(11,8,'imm[4:1]',w,'Младшая группа байтового смещения; bit0=0'),field(7,7,'imm[11]',w,'Переставленный бит смещения'),field(6,0,'opcode',w,'1100011: BRANCH')]
        hint=f'Если {REGS[rs1]} {"=" if name=="beq" else "≠"} {REGS[rs2]}, PC ← 0x{target:08X}; иначе PC ← PC+4. Смещение {imm:+d} байт.'
    elif op==0x23:
        if f3 != 2:raise ValueError('only SW allowed')
        name='sw';fmt='S';imm=sg((((w>>25)&127)<<5)|((w>>7)&31),12)
        asm=f'sw {REGS[rs2]}, {imm}({REGS[rs1]})'
        fs=[field(31,25,'imm[11:5]',w,'Старшие биты смещения со знаком'),field(24,20,'rs2',w,f'Источник записываемых данных: x{rs2} = {REGS[rs2]}'),field(19,15,'rs1',w,f'Базовый адрес: x{rs1} = {REGS[rs1]}'),field(14,12,'funct3',w,'010 = SW: запись 32-битного слова'),field(11,7,'imm[4:0]',w,'Младшие биты смещения'),field(6,0,'opcode',w,'0100011: STORE')]
        hint=f'MEM32[{REGS[rs1]} {imm:+d}] ← {REGS[rs2]}. Число, находящееся в rs2, не кодируется в самой инструкции.'
    elif op==0x03:
        if f3!=2:raise ValueError('only LW allowed')
        name='lw';fmt='I';imm=sg(w>>20,12)
        asm=f'lw {REGS[rd]}, {imm}({REGS[rs1]})'
        fs=[field(31,20,'imm[11:0]',w,'Знаковое 12-битное смещение от базового адреса'),field(19,15,'rs1',w,f'Базовый регистр x{rs1} = {REGS[rs1]}'),field(14,12,'funct3',w,'010 = LW: загрузка 32-битного слова'),field(11,7,'rd',w,f'Приёмник x{rd} = {REGS[rd]}'),field(6,0,'opcode',w,'0000011: LOAD')]
        hint=f'{REGS[rd]} ← MEM32[{REGS[rs1]} {imm:+d}]. FETCH этой инструкции не удостоверяет регистрацию DATA_READ.'
    elif op==0x67:
        if f3!=0:raise ValueError('JALR funct3 must be 0')
        name='jalr';fmt='I';imm=sg(w>>20,12)
        asm=f'jalr {REGS[rd]}, {imm}({REGS[rs1]})'
        fs=[field(31,20,'imm[11:0]',w,'12-битное знаковое смещение'),field(19,15,'rs1',w,f'Базовый регистр перехода x{rs1} = {REGS[rs1]}'),field(14,12,'funct3',w,'000: JALR'),field(11,7,'rd',w,f'Приёмник адреса возврата x{rd} = {REGS[rd]}'),field(6,0,'opcode',w,'1100111: JALR')]
        hint=f'PC ← ({REGS[rs1]} {imm:+d}) & ~1; {REGS[rd]} ← PC+4. При rd=zero значение возврата отбрасывается.'
    elif op==0x13:
        if f3==5:
            if w>>25!=0:raise ValueError('non-SRLI shift in ROM')
            name='srli';fmt='I-shift';imm=(w>>20)&31
            asm=f'srli {REGS[rd]}, {REGS[rs1]}, {imm}'
            fs=[field(31,25,'funct7',w,'0000000 — логический сдвиг вправо'),field(24,20,'shamt[4:0]',w,f'Сдвиг на {imm} разрядов'),field(19,15,'rs1',w,f'Источник x{rs1} = {REGS[rs1]}'),field(14,12,'funct3',w,'101: сдвиг вправо'),field(11,7,'rd',w,f'Приёмник x{rd} = {REGS[rd]}'),field(6,0,'opcode',w,'0010011: OP-IMM')]
            hint=f'{REGS[rd]} ← unsigned({REGS[rs1]}) >> {imm} (заполнение старших разрядов нулями).'
        elif f3 in (0,7):
            name={0:'addi',7:'andi'}[f3];fmt='I';imm=sg(w>>20,12)
            asm=f'{name} {REGS[rd]}, {REGS[rs1]}, {imm}'
            fs=[field(31,20,'imm[11:0]',w,'12-битная константа; в RV32I знаково расширяется'),field(19,15,'rs1',w,f'Исходный регистр x{rs1} = {REGS[rs1]}'),field(14,12,'funct3',w,'000=ADDI, 111=ANDI'),field(11,7,'rd',w,f'Приёмник x{rd} = {REGS[rd]}'),field(6,0,'opcode',w,'0010011: OP-IMM')]
            hint=f'{REGS[rd]} ← {REGS[rs1]} {"+" if name=="addi" else "AND"} ({imm:+d})' + ('; арифметика RV32I по модулю 2³².' if name=='addi' else '; побитовое И с sign_extend(imm).')
        else:raise ValueError('unsupported OP-IMM')
    else:
        raise ValueError(f'Not one of the 10 actual LAB-03 mnemonics: opcode 0x{op:02x} PC=0x{pc:08x}')
    assert sum(f['n'] for f in fs)==32,(name,pc,fs)
    parts=''.join(f['bits'] for f in fs)
    assert parts==f'{w:032b}',(name,pc,parts,f'{w:032b}')
    return dict(word=w,pc=pc,hex=f'0x{w:08X}',pc_hex=f'0x{pc:08X}',name=name,format=fmt,asm=asm,imm=imm,target=target,fields=fs,hint=hint)

F_DOC={
 'opcode':'Основная операция ISA (7 бит): LOAD / STORE / BRANCH / JAL / LUI / OP-IMM.',
 'funct3':'Трёхбитное уточнение операции (например, funct3=010 означает LW/SW в соответствующем opcode).',
 'funct7':'Дополнительная функция (7 бит), здесь отличает логический сдвиг SRLI.',
 'rs1':'Номер первого исходного регистра (не его содержимое).',
 'rs2':'Номер второго исходного регистра или регистра данных для SW.',
 'rd':'Номер регистра-приёмника. rd=zero означает, что результат записи в регистр отбрасывается.',
 'imm':'Непосредственный операнд, собираемый из битов инструкции. Для B/J биты расположены не по порядку.',
 'shamt':'Величина сдвига; для RV32I составляет 5 бит.',
}

def kind_color(field_label):
    k='imm' if field_label.startswith('imm') else 'rs1' if field_label=='rs1' else 'rs2' if field_label=='rs2' else 'rd' if field_label=='rd' else 'shamt' if field_label.startswith('shamt') else 'funct' if field_label.startswith('funct') else 'opcode'
    return COLORS[k]

def set_cell(ws,at,value=None,fill=None,font=None,alignment=None,comment=None):
    c=ws[at]
    if value is not None:c.value=value
    if fill:c.fill=PatternFill('solid',fgColor=fill)
    if font:c.font=font
    if alignment:c.alignment=alignment
    if comment:c.comment=Comment(comment,'LAB-03 / RV32I')
    return c

def merged(ws,from_cell,to_cell,value,fill=WHITE,fg=DARK,size=11,bold=False,wrap=True,alignment='left'):
    ws.merge_cells(f'{from_cell}:{to_cell}')
    c=ws[from_cell];c.value=value
    # Explicit fill across merged region makes LibreOffice stable.
    for row in ws[f'{from_cell}:{to_cell}']:
        for cell in row:cell.fill=PatternFill('solid',fgColor=fill)
    c.font=Font(name='Aptos',size=size,color=fg,bold=bold)
    c.alignment=Alignment(horizontal=alignment,vertical='center',wrap_text=wrap,indent=1 if alignment=='left' else 0)
    return c

def setup_sheet(ws,title,subtitle):
    ws.sheet_view.showGridLines=False
    ws.sheet_view.zoomScale=85
    ws.sheet_properties.pageSetUpPr.fitToPage=True
    ws.page_setup.fitToWidth=1;ws.page_setup.fitToHeight=0
    ws.sheet_view.tabSelected=False
    ws.column_dimensions['A'].width=4
    for idx in range(2,10):ws.column_dimensions[get_column_letter(idx)].width=23
    for idx in range(10,15):ws.column_dimensions[get_column_letter(idx)].width=12
    merged(ws,'A1','N2',title,DARK,WHITE,20,True)
    merged(ws,'A3','N3',subtitle,PANEL,'D6E7F9',10)
    ws.row_dimensions[1].height=29;ws.row_dimensions[2].height=17;ws.row_dimensions[3].height=44
    base=['СТАРТ','НАБЛЮДЕНИЕ','СРАВНЕНИЕ','LED / MMIO','КОМАНДЫ','12 ТРАСС','RAW T3','ROM','ТЕРМИНЫ','ПРОИСХОЖДЕНИЕ','БИТЫ 202','10 КАРТОЧЕК']
    anchors=['ИНСТРУКЦИЯ','ПУЛЬТ','СРАВНЕНИЕ','МОДЕЛЬ MMIO','КОМАНДЫ','ЗАХВАТЫ','ДАННЫЕ T3','ROM LAB03','СЛОВАРЬ','ИСТОЧНИКИ','РАЗБОР RV32I','10 КАРТОЧЕК']
    for i,(text,place) in enumerate(zip(base,anchors),1):
        c=ws.cell(4,i);c.value=text;c.fill=PatternFill('solid',fgColor=TEAL if place==ws.title else GRAY)
        c.font=Font(name='Aptos',size=8,bold=True,color=WHITE);c.alignment=Alignment(horizontal='center',vertical='center')
        c.hyperlink=f"#'{place}'!A1"
    ws.row_dimensions[4].height=28
    ws.sheet_properties.outlinePr.summaryBelow=True

def build(base_xlsx:Path,rom_hex:Path,destination:Path):
    raw=rom_hex.read_bytes();assert hashlib.sha256(raw).hexdigest()==ROM_SHA256, 'Unrecognized LAB-03 ROM SHA256'
    lines=raw.decode('ascii').splitlines();assert len(lines)==256
    romwords=[int(x.strip(),16) for x in lines]
    wb=load_workbook(base_xlsx)
    assert wb['ROM LAB03'].max_row>=261
    old_asm=wb['ROM LAB03'];observations=collections.Counter()
    for row in wb['ДАННЫЕ T3'].iter_rows(min_row=6,max_row=286,values_only=True):
        if row[4]=='FETCH':
            observations[row[5]]+=1
            pc=int(row[5],16)
            assert int(row[6],16)==romwords[pc//4]
    assert sum(observations.values())==269 and len(observations)==29
    assert len({wb['ДАННЫЕ T3'].cell(r,1).value for r in range(6,287)})==12
    records=[]
    for i,w in enumerate(romwords[:202]):
        r=decode(w,i*4);oldpc=old_asm.cell(i+6,2).value;oldword=old_asm.cell(i+6,3).value;oldtxt=old_asm.cell(i+6,5).value
        assert r['pc_hex'].upper()==oldpc.upper() and r['hex'].upper()==oldword.upper(),(i,oldpc,oldword,r)
        assert r['asm'].lower()==oldtxt.lower(),(i,r['asm'],oldtxt)
        r['seen']=observations[r['pc_hex']]
        records.append(r)
    assert all(w==0x13 for w in romwords[202:]), 'Unexpected ROM tail: not 54 padded NOPs'
    assert set(r['name'] for r in records)==set(INFO),'Expected exactly 10 real mnemonics'
    assert set(r['name'] for r in records if r['seen'])=={'addi','andi','beq','jal','lw','sw'}
    assert records[64]['hex']=='0x01D2A023'
    assert [(f['range'],f['bits'],f['label']) for f in records[64]['fields']]==[
        ('31–25','0000000','imm[11:5]'),('24–20','11101','rs2'),('19–15','00101','rs1'),
        ('14–12','010','funct3'),('11–7','00000','imm[4:0]'),('6–0','0100011','opcode')]

    cat=wb.create_sheet('КАТАЛОГ БИТОВ', index=5)
    setup_sheet(cat,'LAB-03 / МАШИННЫЕ КОМАНДЫ — 202 РАЗБОРА',
                'Все 202 слова действующей прошивки (ROM 0x000…0x324). Только 29 PC наблюдались в серии 12 T3. 54 хвостовых NOP — заполнение, не активная программа.')
    merged(cat,'B6','N6','Каталог: каждый адрес PC, машинное слово, инструкция, тип кодирования и поля [31:0] из неизменной LAB-03 ROM.',LIGHT,BLUE,11,True)
    cols=['№','PC','WORD','RV32I ASM','Команда','Формат','FETCH (12)','imm, dec','Target','Binary [31:0]','Действие / замечание','Статус']
    for j,txt in enumerate(cols,1):
        c=cat.cell(8,j,txt);c.fill=PatternFill('solid',fgColor=PANEL);c.font=Font(name='Aptos',size=10,color=WHITE,bold=True);c.alignment=Alignment(horizontal='center',vertical='center',wrap_text=True)
    for k in range(8):
        c=cat.cell(8,k+13,f'Поле {k+1} · биты / значение / имя');c.fill=PatternFill('solid',fgColor=PANEL);c.font=Font(name='Aptos',size=9,color=WHITE,bold=True);c.alignment=Alignment(wrap_text=True,vertical='center')
    cat.row_dimensions[8].height=38
    widths={'A':7,'B':17,'C':16,'D':37,'E':14,'F':12,'G':14,'H':14,'I':17,'J':38,'K':65,'L':21}
    for col,width in widths.items():cat.column_dimensions[col].width=width
    for n in range(13,21):cat.column_dimensions[get_column_letter(n)].width=27
    for i,r in enumerate(records):
        row=9+i
        vals=[i+1,r['pc_hex'],r['hex'],r['asm'],r['name'],r['format'],r['seen'],r['imm'] if r['imm'] is not None else '—',f'0x{r["target"]:08X}' if r['target'] is not None else '—',f'{r["word"]:032b}',r['hint'],'FETCH был' if r['seen'] else 'Только ROM']
        for col,v in enumerate(vals,1):
            c=cat.cell(row,col,v)
            c.font=Font(name='Consolas' if col in (2,3,9,10) else 'Aptos',size=10 if col!=11 else 9,color=GREEN if col in(2,3) else (TEAL if col==7 and r['seen'] else (BLUE if col==1 else DARK)),bold=(col in(3,7) and bool(r['seen'])))
            c.fill=PatternFill('solid',fgColor=('E2F3F0' if r['seen'] and col in (3,7,12) else (WHITE if i%2==0 else LIGHT)))
            c.alignment=Alignment(vertical='center',wrap_text=col in(4,11,12))
        for k,f in enumerate(r['fields']):
            cc=cat.cell(row,13+k,f'{f["range"]}\n{f["bits"]}\n{f["label"]}')
            bg,fg=kind_color(f['label']);cc.fill=PatternFill('solid',fgColor=bg)
            cc.font=Font(name='Consolas',size=9,color=fg,bold=True)
            cc.alignment=Alignment(horizontal='center',vertical='center',wrap_text=True)
            cc.comment=Comment(f['meaning'],'RISC-V / LAB-03')
        cat.row_dimensions[row].height=52
    cat.auto_filter.ref=f'A8:L{8+len(records)}';cat.freeze_panes='D9'
    cat.print_title_rows='1:8'
    # Dynamic decoder's primary navigation is the original T3 inspector; alternate source is any ROM position.
    ws=wb.create_sheet('РАЗБОР RV32I', index=6)
    setup_sheet(ws,'LAB-03 / ДЕКОДЕР RISC-V · 32 БИТА',
                'Переключатель «ТРАССА» привязан к ПУЛЬТУ. «РОМ» даёт доступ ко всем 202 рабочим словам. Поля и биты не придуманы — извлечены из конкретных машинных кодов.')
    merged(ws,'B5','I5','НАСТРОЙКА · выберите источник команды и номер слова программы',PANEL,WHITE,12,True)
    for addr,value in [('B6','№ слова ROM (1…202)'),('F6','Источник'),('B7','Выбранный PC'),('F7','№ по ROM'),('B9','Машинное слово'),('F9','Команда / формат'),('B10','Инструкция ASM'),('B11','Статус по трассам'),('F11','FETCH в 12 RAW')]:
        ws[addr]=value;ws[addr].font=Font(name='Aptos',size=10,color=GRAY,bold=True)
    ws['D6']=65;ws['D6'].font=Font(name='Aptos',size=17,bold=True,color=BLUE);ws['D6'].fill=PatternFill('solid',fgColor=GOLD);ws['D6'].alignment=Alignment(horizontal='center')
    ws['H6']='ТРАССА';ws['H6'].font=Font(name='Aptos',size=14,bold=True,color=BLUE);ws['H6'].fill=PatternFill('solid',fgColor=GOLD);ws['H6'].alignment=Alignment(horizontal='center')
    dv_idx=DataValidation(type='whole',operator='between',formula1='1',formula2='202',allow_blank=False)
    dv_idx.errorTitle='Неверный номер ROM';dv_idx.error='Укажите целое число 1…202';dv_idx.showErrorMessage=True;dv_idx.errorStyle='stop';dv_idx.promptTitle='Рабочие слова ROM';dv_idx.prompt='Индекс 65 соответствует PC=0x00000100';dv_idx.showInputMessage=True
    ws.add_data_validation(dv_idx);dv_idx.add(ws['D6'])
    dv_src=DataValidation(type='list',formula1='"ТРАССА,РОМ"',allow_blank=False)
    dv_src.showErrorMessage=True;dv_src.error='Выберите ТРАССА или РОМ';dv_src.errorTitle='Выбор источника';ws.add_data_validation(dv_src);dv_src.add(ws['H6'])
    # D7 resolves PC, H7 resolves source row: D6 ROM 1-based index / chosen live saved T3 FETCH.
    ws['D7']="=IF($H$6=\"ТРАССА\",IF('ПУЛЬТ'!$D$12=\"FETCH\",'ПУЛЬТ'!$D$13,\"\"),IFERROR(INDEX('КАТАЛОГ БИТОВ'!$B$9:$B$210,$D$6),\"\"))"
    ws['H7']="=IF($D$7=\"\",\"\",IFERROR(MATCH($D$7,'КАТАЛОГ БИТОВ'!$B$9:$B$210,0),\"\"))"
    def getcol(letter,default=''):
        return f'=IF($H$7="","",IFERROR(INDEX(\'КАТАЛОГ БИТОВ\'!${letter}$9:${letter}$210,$H$7),""))'
    ws['D9']=getcol('C');ws['H9']=getcol('E');ws['I9']=getcol('F');ws['D10']=getcol('D');ws['D11']=getcol('L');ws['H11']=getcol('G')
    for a in ['D7','D9','H9','I9','D10','D11','H11','H7']:
        c=ws[a];c.fill=PatternFill('solid',fgColor='E2F3F0');c.font=Font(name='Consolas' if a in ('D7','D9','D10') else 'Aptos',size=11,color=GREEN,bold=a in ('D9','D10'));c.alignment=Alignment(wrap_text=True,vertical='center')
    ws.row_dimensions[6].height=36;ws.row_dimensions[7].height=30;ws.row_dimensions[9].height=32;ws.row_dimensions[10].height=33;ws.row_dimensions[11].height=36
    merged(ws,'B13','I13','ПОБИТОВЫЙ РАЗБОР · точное расположение полей в машинном слове [31:0]',PANEL,WHITE,11,True)
    for k in range(8):
        col=get_column_letter(13+k)
        slot=get_column_letter(2+k)
        ws[f'{slot}15']=getcol(col)
        ws[f'{slot}15'].alignment=Alignment(horizontal='center',vertical='center',wrap_text=True)
        ws[f'{slot}15'].font=Font(name='Consolas',size=12,color='17528C',bold=True)
        ws[f'{slot}15'].fill=PatternFill('solid',fgColor=('DDECF8' if k%2 else 'DDF4F0'))
        ws[f'{slot}15'].comment=Comment('Сверху указан диапазон бит, ниже — их фактические 0/1, затем имя поля ISA. Подробное значение поля см. в каталоге.','LAB-03')
    ws.row_dimensions[15].height=105
    merged(ws,'B17','I17','ПОЛЯ: что означает каждое значение в данной инструкции',PANEL,WHITE,11,True)
    for k in range(8):
        r=19+k
        ws[f'B{r}']=f'{k+1:02d}';ws[f'B{r}'].font=Font(name='Aptos',size=11,bold=True,color=TEAL)
        j=get_column_letter(13+k)
        ws.merge_cells(start_row=r,start_column=3,end_row=r,end_column=4)
        ws[f'C{r}']=getcol(j);ws[f'C{r}'].font=Font(name='Consolas',size=10,color=BLUE)
        ws[f'C{r}'].alignment=Alignment(wrap_text=True,vertical='center')
        # Field meaning is instruction-format-dependent; link to full card is exact source.
        ws.merge_cells(start_row=r,start_column=5,end_row=r,end_column=9)
        ws[f'E{r}']=f'=IF($H$7="","",IFERROR(INDEX(\'КАТАЛОГ ОПИСАНИЙ\'!${get_column_letter(2+k)}$9:${get_column_letter(2+k)}$210,$H$7),""))'
        ws[f'E{r}'].font=Font(name='Aptos',size=10,color=DARK)
        ws[f'E{r}'].alignment=Alignment(wrap_text=True,vertical='center')
        if k%2==0:
            for j2 in range(2,10):ws.cell(r,j2).fill=PatternFill('solid',fgColor=LIGHT)
        ws.row_dimensions[r].height=46
    merged(ws,'B28','I28','ОБЪЯСНЕНИЕ ИНСТРУКЦИИ В НАШЕЙ ПРОГРАММЕ',PANEL,WHITE,11,True)
    ws.merge_cells('B29:I30')
    ws['B29']=getcol('K');ws['B29'].font=Font(name='Aptos',size=12,color=DARK,bold=True)
    ws['B29'].alignment=Alignment(wrap_text=True,vertical='center',indent=1)
    ws['B29'].fill=PatternFill('solid',fgColor='E2F3F0')
    ws.row_dimensions[29].height=47;ws.row_dimensions[30].height=35
    ws['B32']='ИНЖЕНЕРНАЯ ГРАНИЦА';ws['B32'].font=Font(name='Aptos',size=12,bold=True,color='B96C1B')
    merged(ws,'B33','I34','FETCH означает принятый запрос слова из ROM — НЕ доказательство instruction retirement. Значения регистров при FETCH не измеряются. Если выбран LED_WRITE, это событие MMIO, а не инструкция: вернитесь на FETCH.',GOLD,'664614',10)
    ws.row_dimensions[33].height=35;ws.row_dimensions[34].height=43
    ws.freeze_panes='B5';ws.sheet_view.zoomScale=85

    # Exact per-field semantic explanation, aligned with catalogue row and format.
    explain=wb.create_sheet('КАТАЛОГ ОПИСАНИЙ')
    explain.sheet_state='hidden'
    for k in range(8):explain.cell(8,k+2,f'field{k+1}')
    for i,r in enumerate(records):
        for k,f in enumerate(r['fields']):explain.cell(9+i,2+k,f['meaning'])

    cards=wb.create_sheet('10 КАРТОЧЕК', index=7)
    setup_sheet(cards,'LAB-03 / 10 МНЕМОНИК НА РЕАЛЬНЫХ МАШИННЫХ СЛОВАХ',
                'Один подлинный пример каждой встречающейся в LAB-03 мнемоники. 6 встречались среди 269 FETCH, 4 имеются только в активной ROM. R-type в этой программе нет.')
    rep_order=['sw','lw','addi','andi','beq','jal','lui','bne','srli','jalr']
    source=wb['КОМАНДЫ']; rep={source.cell(row,1).value:source.cell(row,10).value for row in range(7,17)}
    seenct=collections.Counter()
    for i,name in enumerate(rep_order):
        address=rep[name];r=records[int(address,16)//4];assert r['name']==name,(name,r)
        start=7+i*11
        tag='НАБЛЮДАЛОСЬ НА FPGA' if r['seen'] else 'ПРИСУТСТВУЕТ В ROM (в серии T3 не наблюдалось)'
        merged(cards,f'B{start}',f'I{start}',f'{i+1:02d}  {name.upper()}  •  {INFO[name][0]}   |   PC={r["pc_hex"]}   |   {tag}',TEAL if r['seen'] else PURPLE,WHITE,11,True)
        cards.row_dimensions[start].height=32
        merged(cards,f'B{start+1}',f'I{start+1}',f'{r["hex"]}     {r["asm"]}     ·     {r["format"]}-type / {r["format"]}',LIGHT,BLUE,15,True)
        cards.row_dimensions[start+1].height=34
        merged(cards,f'B{start+2}',f'I{start+2}',INFO[name][2],WHITE,DARK,10)
        cards.row_dimensions[start+2].height=35
        for k in range(8):
            col=get_column_letter(2+k)
            if k<len(r['fields']):
                f=r['fields'][k];bg,fg=kind_color(f['label'])
                for off,value in [(3,f['range']),(4,f['bits']),(5,f['label'])]:
                    c=cards[f'{col}{start+off}'];c.value=value;c.fill=PatternFill('solid',fgColor=bg)
                    c.font=Font(name='Consolas' if off==4 else 'Aptos',size=11 if off==4 else 9,color=fg,bold=off in (4,5));c.alignment=Alignment(horizontal='center',vertical='center',wrap_text=True)
                    c.comment=Comment(f['meaning'],'RISC-V / LAB-03')
            else:
                for off in (3,4,5):cards[f'{col}{start+off}'].fill=PatternFill('solid',fgColor=WHITE)
        for off in (3,4,5):cards.row_dimensions[start+off].height=29 if off!=4 else 40
        merged(cards,f'B{start+6}',f'I{start+6}',f'{r["format"]}-format / смещение: {r["imm"] if r["imm"] is not None else "см. U immediate"}   |   {r["hint"]}', 'EAF3F9',DARK,10)
        cards.row_dimensions[start+6].height=58
        merged(cards,f'B{start+7}',f'I{start+7}',f'Роль в LAB-03: {INFO[name][3]}  |  FETCH этого PC в архиве: {r["seen"]}',WHITE,GRAY,10)
        cards.row_dimensions[start+7].height=50
        cards.row_dimensions[start+8].height=12
        seenct[name]=r['seen']
    cards.freeze_panes='B5';cards.print_options.horizontalCentered=True
    cards.print_area=f'A1:N{start+9}'
    # Append navigators without touching previous calculations.
    for sheet in wb:
        if sheet.title in ('КАТАЛОГ ОПИСАНИЙ',):continue
        if sheet.title in ('КАТАЛОГ БИТОВ','РАЗБОР RV32I','10 КАРТОЧЕК'):continue
        for col,txt,where in [('K','БИТЫ 202','РАЗБОР RV32I'),('L','10 КАРТОЧЕК','10 КАРТОЧЕК')]:
            c=sheet[f'{col}4'];c.value=txt;c.hyperlink=f"#'{where}'!A1";c.font=Font(name='Aptos',size=8,bold=True,color=WHITE)
            c.fill=PatternFill('solid',fgColor=PURPLE);c.alignment=Alignment(horizontal='center',vertical='center')
    main=wb['ИНСТРУКЦИЯ'];main['B14']='07 · БИТОВЫЙ ДЕКОДЕР'
    main['E14']='Новый лист «РАЗБОР RV32I»: источник ТРАССА автоматически повторяет FETCH из ПУЛЬТА. В режиме РОМ индекс 1…202 выбирает любое рабочее слово.'
    main['B15']='08 · КАРТОЧКИ'
    main['E15']='Лист «10 КАРТОЧЕК» подробно показывает битовые поля каждого типа команды на настоящем машинном слове; «КАТАЛОГ БИТОВ» раскрывает все 202 адреса.'
    for row in (14,15):
        main[f'B{row}'].font=Font(name='Aptos',size=11,bold=True,color=TEAL)
        main[f'E{row}'].font=Font(name='Aptos',size=10,color=DARK)
        main[f'E{row}'].alignment=Alignment(wrap_text=True,vertical='center')
        main.row_dimensions[row].height=37
    # Source note as self-contained workbook cross-reference.
    sources=wb['ИСТОЧНИКИ'];sources['B36']='Битовый декодер RV32I';sources['E36']='202 слова из LAB-03 ROM; FORMAT I / I-shift / S / B / U / J. Только 29 адресов наблюдались в 12 T3.'
    sources['B37']='ROM SHA256';sources['E37']=ROM_SHA256
    sources['E36'].alignment=Alignment(wrap_text=True);sources.row_dimensions[36].height=41
    # Workbook calculation setting: LibreOffice Calc recomputes formulas on change.
    wb.calculation=CalcProperties(calcMode='auto',fullCalcOnLoad=True,forceFullCalc=True)
    destination.parent.mkdir(parents=True,exist_ok=True)
    wb.active=wb.sheetnames.index('РАЗБОР RV32I')
    wb.save(destination)
    restore_cached_formula_values(base_xlsx,destination,initial=records[64])
    # Basic content and formula audit.
    test=load_workbook(destination,data_only=False,read_only=False)
    assert test['РАЗБОР RV32I']['D6'].value==65
    assert test['РАЗБОР RV32I']['H6'].value=='ТРАССА'
    assert len(records)==202 and len(seenct)==10
    assert 'РОМ' in str(test['РАЗБОР RV32I'].data_validations.dataValidation)
    assert test['КАТАЛОГ БИТОВ']['C73'].value=='0x01D2A023'
    for i,r in enumerate(records):
        assert test['КАТАЛОГ БИТОВ'].cell(i+9,3).value==r['hex']
    return records

OO='http://schemas.openxmlformats.org/spreadsheetml/2006/main'
NS={'m':OO}

def restore_cached_formula_values(oldbook:Path,newbook:Path,initial:dict):
    """Retain calculation caches of existing 1741 formulas, seed the new demo view.

    This patches only OOXML <v> caches and correct string t attrs, never the formula.
    LibreOffice Calc independently recomputes all formula cells after user edits.
    """
    tmp=newbook.with_suffix('.cached.tmp')
    with ZipFile(oldbook,'r') as zin, ZipFile(newbook,'r') as zout, ZipFile(tmp,'w',ZIP_DEFLATED) as znew:
        oldparts={name:zin.read(name) for name in zin.namelist() if re.fullmatch(r'xl/worksheets/sheet\d+\.xml',name)}
        cached=0
        # The first 10 existing worksheets keep their original order? New sheets inserted
        # at index 5, so map by Excel workbook relationships/title, not naïve sheet numbers.
        def workbook_map(z):
            root=ET.fromstring(z.read('xl/workbook.xml'))
            relns={'r':'http://schemas.openxmlformats.org/officeDocument/2006/relationships'}
            relroot=ET.fromstring(z.read('xl/_rels/workbook.xml.rels'))
            rel={e.attrib['Id']:e.attrib['Target'].split('/')[-1] for e in relroot}
            return {x.attrib['name']:'xl/worksheets/'+rel[x.attrib['{'+relns['r']+'}id']] for x in root.findall('m:sheets/m:sheet',NS)}
        oldmap=workbook_map(zin);newmap=workbook_map(zout)
        reverse={v:k for k,v in newmap.items()}
        cached_new={}
        def new_val(at,value):
            cached_new[at]=value
        # Defaults: PULT capture 5 / event 27 is SW. All formulas below should reflect its data.
        new_val('D7','0x00000100');new_val('H7',65);new_val('D9',initial['hex']);new_val('H9',initial['name']);new_val('I9',initial['format']);new_val('D10',initial['asm']);new_val('D11','FETCH был');new_val('H11',initial['seen']);new_val('B29',initial['hint'])
        for k in range(8):
            f=initial['fields'][k] if k<len(initial['fields']) else None
            new_val(f'{get_column_letter(2+k)}15', f'{f["range"]}\n{f["bits"]}\n{f["label"]}' if f else '')
            new_val(f'C{19+k}',f'{f["range"]}\n{f["bits"]}\n{f["label"]}' if f else '')
            new_val(f'E{19+k}',f['meaning'] if f else '')
        oldvalues={}
        for name,oldxml in oldparts.items():
            root=ET.fromstring(oldxml)
            these={}
            for c in root.findall('.//m:c',NS):
                if c.find('m:f',NS) is not None and c.find('m:v',NS) is not None:
                    these[c.attrib['r']]=(c.find('m:v',NS).text,c.attrib.get('t'))
            oldvalues[name]=these
        for item in zout.infolist():
            fname=item.filename
            if fname not in reverse:
                znew.writestr(item,zout.read(fname));continue
            sheetname=reverse[fname]
            root=ET.fromstring(zout.read(fname))
            restored=0
            if sheetname in oldmap:
                old_data=oldvalues.get(oldmap[sheetname],{})
                for c in root.findall('.//m:c',NS):
                    if c.find('m:f',NS) is None:continue
                    orig=old_data.get(c.attrib['r'])
                    if not orig:continue
                    value, t=orig
                    if t is not None:c.attrib['t']=t
                    v=c.find('m:v',NS)
                    if v is None:v=ET.SubElement(c,'{%s}v'%OO)
                    v.text=value
                    restored+=1
            elif sheetname=='РАЗБОР RV32I':
                for c in root.findall('.//m:c',NS):
                    if c.find('m:f',NS) is None:continue
                    ref=c.attrib['r']
                    if ref not in cached_new:continue
                    value=cached_new[ref]
                    c.attrib['t']='str' if isinstance(value,str) else 'n'
                    v=c.find('m:v',NS)
                    if v is None:v=ET.SubElement(c,'{%s}v'%OO)
                    v.text=str(value)
                    restored+=1
            cached+=restored
            znew.writestr(item, ET.tostring(root,encoding='utf-8',xml_declaration=False))
    shutil.move(tmp,newbook)
    assert cached>1700,('Expected to preserve existing formula caches',cached)
    # Check actual cached value entry for new decoder.
    w=load_workbook(newbook,data_only=True)
    assert w['РАЗБОР RV32I']['D9'].value=='0x01D2A023',(cached,w['РАЗБОР RV32I']['D9'].value)
    assert w['РАЗБОР RV32I']['B15'].value=='31–25\n0000000\nimm[11:5]'
    assert w['РАЗБОР RV32I']['G15'].value=='6–0\n0100011\nopcode'
    print('LAB03_ISA_CACHE_PASS: cached_formulas=',cached)


def main():
    p=argparse.ArgumentParser(description=__doc__)
    p.add_argument('--base',type=Path,default=Path(__file__).with_name('LAB03_PicoRV32_Excel_Observatory_v1.0.xlsx'))
    p.add_argument('--rom',type=Path,default=Path(__file__).with_name('lab03_boot.hex'))
    p.add_argument('--out',type=Path,default=Path(__file__).with_name('LAB03_PicoRV32_Excel_Observatory_v1.1_RV32I.xlsx'))
    a=p.parse_args()
    rec=build(a.base,a.rom,a.out)
    print('LAB03_ISA_ALL_ROM_DECODE_PASS: records=202 mnemonics=10 formats=I,I-shift,S,B,U,J')
    print('LAB03_ISA_OBSERVED_PASS: raw_fetch=269 distinct_PC=29 trace_mnemonics=6')
    print('LAB03_ISA_SAMPLE_PASS: pc=0x00000100 word=0x01D2A023 asm=sw t4, 0(t0)')
    print('LAB03_ISA_OUT:',a.out, 'size=',a.out.stat().st_size)

if __name__=='__main__':main()
